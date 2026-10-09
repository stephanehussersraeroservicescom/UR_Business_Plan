"""
Execute le modele et produit quatre sorties dans out/ :
  rapport.html       presentation cliquable, chaque chiffre remonte a ses preuves
  resultats.json     toutes les grandeurs calculees
  tracabilite.md     version texte de la chaine de preuve
  compte-resultat.xlsx

Usage :
    python run.py
    python run.py --set nb_unites=5 --set gate_fee=250
    python run.py --set part_dette=0 --set taux_interet=0 --rembourse 3
"""
import argparse, json, sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))
from engine.modele import executer  # noqa: E402
from engine.rapport import generer  # noqa: E402

RACINE = Path(__file__).resolve().parent
OUT = RACINE / "out"
OUT.mkdir(exist_ok=True)

ORDRE_FIABILITE = ["certifie_tiers", "mesure_interne", "constructeur", "litterature", "aucune"]


def convertir(v):
    if v in ("true", "false"):
        return v == "true"
    try:
        return int(v) if "." not in v and "e" not in v.lower() else float(v)
    except ValueError:
        try:
            return float(v)
        except ValueError:
            return v


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--set", action="append", default=[], metavar="CLE=VALEUR")
    ap.add_argument("--rembourse", type=int, default=None,
                    help="annees d'exploitation pour rembourser 100%% du CAPEX")
    a = ap.parse_args()
    surcharges = {k: convertir(v) for k, v in (s.split("=", 1) for s in a.set)}

    reg, m = executer(surcharges)

    pbs = reg.controle()
    print(f"Controle du registre : {len(pbs)} signalement(s)")
    for p in pbs[:20]:
        print("   -", p)

    res = {k: v["valeur"] for k, v in m.r.items()}
    if a.rembourse:
        res["gate_fee_requis"] = m.gate_fee_requis(a.rembourse)
        res["_annees_remboursement"] = a.rembourse
    (OUT / "resultats.json").write_text(
        json.dumps({"resultats": res, "compte_resultat": m.pl,
                    "surcharges": surcharges, "controle": pbs},
                   indent=2, ensure_ascii=False, default=str), encoding="utf-8")

    generer(reg, m, OUT / "rapport.html", surcharges)

    li = ["# Tracabilite du modele Urban Rig", ""]
    if surcharges:
        li += ["## Surcharges appliquees", ""] + [f"- `{k}` = {v}" for k, v in surcharges.items()] + [""]
    li += ["## Resultats et chaine de preuve", ""]
    for cle, d in m.r.items():
        v = d["valeur"]
        li.append(f"### {cle}")
        li.append(f"**{v:,.4g} {d['unite']}**" if isinstance(v, float) else f"**{v} {d['unite']}**")
        li.append("")
        if not d["depend"]:
            li += ["_Grandeur derivee d'autres resultats._", ""]
            continue
        li += ["| Parametre | Valeur | Statut | Decision | Preuves | Fiabilite |",
               "|---|---:|---|---|---|---|"]
        for k in sorted(d["depend"]):
            pr = reg.params[k]
            preuves = pr.get("preuves", [])
            fiab = min((ORDRE_FIABILITE.index(reg.preuves[p]["fiabilite"])
                        for p in preuves if p in reg.preuves), default=len(ORDRE_FIABILITE) - 1)
            li.append(f"| {pr['label']} | {pr['valeur']} {pr.get('unite','')} | {pr['statut']} "
                      f"| {pr.get('decision','-')} | {', '.join(preuves) or '-'} "
                      f"| {ORDRE_FIABILITE[fiab]} |")
        li.append("")
    li += ["## Journal des decisions", ""]
    for did, d in reg.decisions.items():
        li += [f"### {did} — {d['titre']}", f"_{d['date']}_", "",
               f"- **De** : {d.get('de','-')} → **vers** : {d.get('vers','-')} {d.get('unite','')}",
               f"- **Motif** : {d['motif'].strip()}",
               f"- **Preuves** : {', '.join(d.get('preuves', []))}"]
        if d.get("a_verifier"):
            li.append(f"- **Reste a verifier** : {d['a_verifier']}")
        li.append("")
    (OUT / "tracabilite.md").write_text("\n".join(li), encoding="utf-8")

    try:
        import openpyxl
        from openpyxl.styles import Font, PatternFill
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Compte de resultat"
        cols = list(m.pl[0].keys())
        for i, c in enumerate(cols, 1):
            cell = ws.cell(row=1, column=i, value=c)
            cell.font = Font(bold=True)
            cell.fill = PatternFill("solid", fgColor="D9E2F3")
        for r, ligne in enumerate(m.pl, 2):
            for i, c in enumerate(cols, 1):
                ws.cell(row=r, column=i,
                        value=round(ligne[c], 2) if isinstance(ligne[c], float) else ligne[c])
        ws2 = wb.create_sheet("Parametres")
        for i, h in enumerate(["Cle", "Domaine", "Label", "Valeur", "Unite", "Statut",
                               "Decision", "Preuves", "Note"], 1):
            cell = ws2.cell(row=1, column=i, value=h)
            cell.font = Font(bold=True)
            cell.fill = PatternFill("solid", fgColor="D9E2F3")
        for r, (k, pr) in enumerate(sorted(reg.params.items(),
                                           key=lambda x: (x[1]["_domaine"], x[0])), 2):
            for i, v in enumerate([k, pr["_domaine"], pr["label"], pr["valeur"], pr.get("unite", ""),
                                   pr["statut"], pr.get("decision", ""),
                                   ", ".join(pr.get("preuves", [])), pr.get("note", "")], 1):
                ws2.cell(row=r, column=i, value=v)
        ws3 = wb.create_sheet("Resultats")
        ws3.append(["Grandeur", "Valeur", "Unite"])
        for c in "ABC":
            ws3[f"{c}1"].font = Font(bold=True)
        for k, d in m.r.items():
            ws3.append([k, round(d["valeur"], 4) if isinstance(d["valeur"], float) else d["valeur"],
                        d["unite"]])
        for w in (ws, ws2, ws3):
            for col in w.columns:
                w.column_dimensions[col[0].column_letter].width = 22
        wb.save(OUT / "compte-resultat.xlsx")
    except ImportError:
        print("openpyxl absent : export xlsx ignore")

    print()
    for k in ["tonnage_annuel", "couverture_thermique", "fioul_litres_jour", "capex_total",
              "capex_par_tonne_an", "ebitda_an3", "marge_ebitda_an3", "gate_fee_equilibre",
              "intensite_carbone_site", "ecart_bouclage_carbone"]:
        d = m.r[k]
        print(f"{k:<28} {d['valeur']:>16,.2f}  {d['unite']}")
    if a.rembourse:
        print(f"{'gate_fee_requis':<28} {res['gate_fee_requis']:>16,.2f}  EUR/T "
              f"(remboursement en {a.rembourse} ans)")
    print(f"\nSorties ecrites dans {OUT}")
    print(f"  -> ouvrir {OUT / 'rapport.html'} dans un navigateur")


if __name__ == "__main__":
    main()
